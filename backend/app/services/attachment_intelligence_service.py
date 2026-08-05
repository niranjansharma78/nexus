from __future__ import annotations

import csv
import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".xlsx",
    ".xlsm",
    ".csv",
    ".docx",
    ".txt",
}

PO_HINTS = {
    "purchase order",
    "po no",
    "po number",
    "buyer",
    "delivery date",
}

DISPATCH_HINTS = {
    "dispatch",
    "despatch",
    "invoice no",
    "lr no",
    "transporter",
}

VALUE_PATTERNS = [
    re.compile(
        r"(?:grand total|total value|order value|invoice value|"
        r"net amount|total amount)\s*[:\-]?\s*"
        r"(?:₹|rs\.?|inr)?\s*([0-9][0-9,]*(?:\.[0-9]{1,2})?)",
        re.IGNORECASE,
    ),
    re.compile(
        r"(?:₹|rs\.?|inr)\s*([0-9][0-9,]*(?:\.[0-9]{1,2})?)",
        re.IGNORECASE,
    ),
]

KM_PATTERNS = [
    re.compile(
        r"(?:total\s+)?(?:quantity|qty|length)\s*[:\-]?\s*"
        r"([0-9][0-9,]*(?:\.[0-9]+)?)\s*"
        r"(?:km|kms|kilometres|kilometers)",
        re.IGNORECASE,
    ),
    re.compile(
        r"\b([0-9][0-9,]*(?:\.[0-9]+)?)\s*"
        r"(?:km|kms|kilometres|kilometers)\b",
        re.IGNORECASE,
    ),
]

PO_NUMBER_PATTERNS = [
    re.compile(
        r"\b(?:po|purchase\s+order)\s*"
        r"(?:number|no\.?|#)\s*[:\-]?\s*"
        r"([A-Z0-9][A-Z0-9/_-]{3,})\b",
        re.IGNORECASE,
    ),
    re.compile(
        r"(?im)^\s*po\s*[:\-]\s*"
        r"([A-Z0-9][A-Z0-9/_-]{3,})\s*$",
        re.IGNORECASE,
    ),
]

INVOICE_PATTERNS = [
    re.compile(
        r"\b(?:invoice|inv)\s*(?:number|no\.?|#)?\s*[:\-]?\s*"
        r"([A-Z0-9/_-]{4,})",
        re.IGNORECASE,
    ),
]

LR_PATTERNS = [
    re.compile(
        r"\b(?:lr|awb|consignment)\s*(?:number|no\.?|#)?\s*[:\-]?\s*"
        r"([A-Z0-9/_-]{4,})",
        re.IGNORECASE,
    ),
]

CUSTOMER_PATTERNS = [
    re.compile(
        r"\b(?:customer|buyer|sold to|ship to|consignee)\s*[:\-]?\s*"
        r"([A-Za-z0-9&.,() /_-]{3,80})",
        re.IGNORECASE,
    ),
]

TRANSPORTER_PATTERNS = [
    re.compile(
        r"\b(?:transporter|carrier|logistics)\s*[:\-]?\s*"
        r"([A-Za-z0-9&.,() /_-]{3,80})",
        re.IGNORECASE,
    ),
]


@dataclass(slots=True)
class AttachmentExtraction:
    document_type: str
    confidence: float
    review_required: bool
    text: str
    structured_data: dict[str, Any]

    def to_dict(self) -> dict:
        return asdict(self)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )
    output: list[str] = []

    for sheet in workbook.worksheets:
        output.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                output.append(" | ".join(values))

    return "\n".join(output)


def _read_csv(path: Path) -> str:
    output: list[str] = []

    with path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline="",
    ) as handle:
        for row in csv.reader(handle):
            values = [str(value).strip() for value in row]
            if any(values):
                output.append(" | ".join(values))

    return "\n".join(output)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    output = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                output.append(" | ".join(values))

    return "\n".join(output)


def extract_text(path: str | Path) -> str:
    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(file_path)

    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported attachment type: {extension}")

    if extension == ".pdf":
        return _read_pdf(file_path)
    if extension in {".xlsx", ".xlsm"}:
        return _read_xlsx(file_path)
    if extension == ".csv":
        return _read_csv(file_path)
    if extension == ".docx":
        return _read_docx(file_path)

    return file_path.read_text(encoding="utf-8", errors="replace")


def _first_group(
    patterns: list[re.Pattern[str]],
    text: str,
) -> str | None:
    for pattern in patterns:
        match = pattern.search(text)
        if match:
            return match.group(1).strip(" .,:;-")
    return None


def _number(
    patterns: list[re.Pattern[str]],
    text: str,
) -> float | None:
    value = _first_group(patterns, text)

    if value is None:
        return None

    try:
        return float(value.replace(",", ""))
    except ValueError:
        return None


def classify_document(text: str) -> tuple[str, float]:
    lowered = text.lower()
    po_score = sum(hint in lowered for hint in PO_HINTS)
    dispatch_score = sum(hint in lowered for hint in DISPATCH_HINTS)

    if po_score >= 2 and po_score >= dispatch_score:
        return "purchase_order", min(0.95, 0.65 + po_score * 0.05)

    if dispatch_score >= 2:
        return "dispatch", min(0.95, 0.65 + dispatch_score * 0.05)

    return "generic", 0.45


def extract_structured_fields(
    text: str,
    document_type: str,
) -> dict[str, Any]:
    value_amount = _number(VALUE_PATTERNS, text)

    data: dict[str, Any] = {
        "customer": _first_group(CUSTOMER_PATTERNS, text),
        "quantity_km": _number(KM_PATTERNS, text),
        "value_amount": value_amount,
        "currency": "INR" if value_amount is not None else None,
    }

    if document_type == "purchase_order":
        data["po_number"] = _first_group(PO_NUMBER_PATTERNS, text)

    if document_type == "dispatch":
        data["invoice_number"] = _first_group(INVOICE_PATTERNS, text)
        data["lr_number"] = _first_group(LR_PATTERNS, text)
        data["transporter"] = _first_group(TRANSPORTER_PATTERNS, text)

    return data


def analyze_attachment(path: str | Path) -> AttachmentExtraction:
    text = extract_text(path)
    cleaned = re.sub(r"\s+", " ", text).strip()
    document_type, confidence = classify_document(cleaned)
    structured_data = extract_structured_fields(cleaned, document_type)

    important_fields = [
        structured_data.get("customer"),
        structured_data.get("quantity_km"),
        structured_data.get("value_amount"),
    ]
    populated = sum(value not in (None, "") for value in important_fields)

    return AttachmentExtraction(
        document_type=document_type,
        confidence=confidence,
        review_required=confidence < 0.75 or populated < 2,
        text=cleaned,
        structured_data=structured_data,
    )
